#!/usr/bin/env python3
"""Genera las plantillas Excel y CSV de la guía en la carpeta plantillas/.

Plantillas:
  - inventario-sistemas-ia.xlsx          Inventario de sistemas de IA
  - matriz-riesgos-ia.xlsx               Metodología, registro de riesgos y mapa de calor
  - declaracion-de-aplicabilidad.xlsx    Declaración de Aplicabilidad (SoA) de los 38 controles
  - declaracion-de-aplicabilidad.csv     La misma SoA en CSV (sin fórmulas)
  - registro-incidentes-ia.xlsx          Registro de incidentes de IA

Los datos de los controles salen de data/controles.yml. Los nombres de los controles
son traducciones libres del autor.

Uso:
    python scripts/generar_plantillas_excel.py
"""
from __future__ import annotations

import csv
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "plantillas"
DATOS = RAIZ / "data" / "controles.yml"

AUTOR = "Carlos Adrián Guzmán · Descifrando ISO 42001"
LICENCIA = "CC BY-SA 4.0 · https://adriangzmncrz-arch.github.io/descifrando-iso42001/"
AVISO = ("Plantilla de apoyo de la guía Descifrando ISO 42001. Interpretación del autor; no sustituye a la norma "
         "ISO/IEC 42001 ni es asesoría legal. Los nombres de los controles son traducciones libres.")

INDIGO = "312E81"
INDIGO_CLARO = "E0E7FF"
CIAN = "0891B2"
GRIS = "F4F5FB"
BORDE = Side(style="thin", color="D9DCEF")

NIVEL_COLOR = {"Bajo": "BBF7D0", "Medio": "FEF08A", "Alto": "FED7AA", "Crítico": "FECACA"}
ROLES = {"usa": "Usa IA de terceros", "desarrolla": "Desarrolla IA", "provee": "Provee IA a clientes"}
NOVEDAD = {"nuevo": "Nuevo", "similar": "Similar", "equivalente": "Equivalente"}


# ----------------------------------------------------------------------------- utilidades

def libro(titulo: str) -> Workbook:
    wb = Workbook()
    wb.properties.creator = AUTOR
    wb.properties.title = titulo
    wb.properties.description = AVISO
    wb.properties.keywords = "ISO/IEC 42001; SGIA; plantilla"
    return wb


def encabezados(ws, fila: int, columnas: list[tuple[str, int]]) -> None:
    for i, (texto, ancho) in enumerate(columnas, 1):
        c = ws.cell(row=fila, column=i, value=texto)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INDIGO)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 36


def titulo_hoja(ws, texto: str, subtitulo: str = "") -> None:
    ws["A1"] = texto
    ws["A1"].font = Font(bold=True, size=14, color=INDIGO)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True, size=9, color="555555")


def hoja_instrucciones(wb: Workbook, titulo: str, pasos: list[str]) -> None:
    ws = wb.active
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 110
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=16, color=INDIGO)
    ws["A2"] = AVISO
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 30
    fila = 4
    for paso in pasos:
        ws.cell(row=fila, column=1, value=paso).alignment = Alignment(wrap_text=True, vertical="top")
        fila += 1
    fila += 1
    ws.cell(row=fila, column=1, value=f"Autor: {AUTOR} · Licencia {LICENCIA}").font = Font(size=9, color="555555")


def hoja_listas(wb: Workbook, listas: dict[str, list[str]]):
    ws = wb.create_sheet("Listas")
    referencias = {}
    for col, (nombre, valores) in enumerate(listas.items(), 1):
        letra = get_column_letter(col)
        ws.cell(row=1, column=col, value=nombre).font = Font(bold=True)
        for fila, valor in enumerate(valores, 2):
            ws.cell(row=fila, column=col, value=valor)
        ws.column_dimensions[letra].width = max(18, max(len(v) for v in valores) + 2)
        referencias[nombre] = f"Listas!${letra}$2:${letra}${len(valores) + 1}"
    ws.sheet_state = "visible"
    return referencias


def validacion_lista(ws, rango: str, referencia: str, mensaje: str = "") -> None:
    dv = DataValidation(type="list", formula1=f"={referencia}", allow_blank=True)
    dv.error = "Elige un valor de la lista."
    dv.errorTitle = "Valor no válido"
    if mensaje:
        dv.prompt = mensaje
        dv.promptTitle = "Ayuda"
    ws.add_data_validation(dv)
    dv.add(rango)


def tabla(ws, nombre: str, rango: str) -> None:
    t = Table(displayName=nombre, ref=rango)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def bordes(ws, rango: str) -> None:
    for fila in ws[rango]:
        for c in fila:
            c.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
            c.alignment = Alignment(wrap_text=True, vertical="top")


# ----------------------------------------------------------------------------- inventario

def inventario() -> Path:
    wb = libro("Inventario de sistemas de IA")
    hoja_instrucciones(wb, "Inventario de sistemas de IA", [
        "1. Registra cada sistema de IA que tu organización usa, desarrolla o provee, incluidos asistentes de IA generativa, "
        "chatbots de proveedores y módulos con IA dentro de otros productos.",
        "2. Determina tu rol frente a cada sistema (cláusula 4.1): un mismo sistema puede implicar varios roles.",
        "3. Las columnas en gris se calculan solas: indican si conviene una evaluación de impacto (6.1.4 y A.5) y la fecha "
        "de la siguiente revisión (12 meses después de la última evaluación).",
        "4. Usa el inventario para fijar el alcance del SGIA (4.3), documentar recursos (A.4) y planificar evaluaciones de "
        "riesgo e impacto (8.2 y 8.4).",
        "5. Las filas de ejemplo corresponden a la empresa ficticia Contadores Alameda (caso práctico 1). Bórralas o "
        "reemplázalas por tus sistemas.",
    ])
    listas = hoja_listas(wb, {
        "Rol": ["Cliente o usuario de IA", "Productor de IA", "Proveedor de IA", "Socio de IA", "Varios roles"],
        "Tipo": ["IA generativa (LLM)", "Aprendizaje automático predictivo", "Visión por computadora",
                 "Procesamiento de voz", "Sistema de recomendación", "Reglas con componentes de IA", "Otro"],
        "SiNo": ["Sí", "No", "No sé"],
        "Nivel": ["Bajo", "Medio", "Alto", "Crítico", "Sin evaluar"],
        "Estado": ["Propuesto", "En evaluación", "Piloto", "En producción", "Suspendido", "Retirado"],
        "Origen": ["Desarrollo propio", "Proveedor externo (SaaS o API)", "Módulo dentro de otro producto",
                   "Código abierto adaptado", "IA no autorizada detectada"],
    })
    ws = wb.create_sheet("Inventario", 1)
    titulo_hoja(ws, "Inventario de sistemas de IA", "Controles relacionados: 4.1, 4.3, A.4.2 a A.4.6, A.5.2, A.9.4")
    cols = [
        ("ID", 8), ("Sistema de IA", 26), ("Propósito y uso previsto", 40), ("Rol de la organización", 20),
        ("Origen", 22), ("Proveedor o equipo responsable", 22), ("Tipo de IA", 22), ("Datos que usa", 30),
        ("¿Datos personales?", 12), ("¿Datos sensibles?", 12), ("¿Influye en decisiones sobre personas?", 16),
        ("Personas o grupos afectados", 26), ("Nivel de riesgo", 12), ("¿Requiere evaluación de impacto?", 16),
        ("Dueño del sistema", 20), ("Estado", 14), ("Última evaluación de riesgos", 14), ("Próxima revisión", 14),
        ("Notas", 30),
    ]
    encabezados(ws, 4, cols)
    ejemplos = [
        ["IA-01", "Asistente de IA generativa de la suite de oficina", "Redactar correos, resumir juntas y analizar hojas de cálculo internas",
         "Cliente o usuario de IA", "Proveedor externo (SaaS o API)", "Proveedor de nube (licencia empresarial)", "IA generativa (LLM)",
         "Correos, documentos internos, hojas de nómina", "Sí", "No", "No", "Colaboradores; clientes cuyos datos aparecen en documentos",
         "Medio", None, "Gerente de TI", "En producción", "2026-03-15", None, "Política de uso aceptable vigente"],
        ["IA-02", "Alma, chatbot de WhatsApp", "Responder preguntas frecuentes de clientes sobre fechas fiscales, estatus de trámites y citas",
         "Cliente o usuario de IA", "Proveedor externo (SaaS o API)", "BotNorte (proveedor ficticio)", "IA generativa (LLM)",
         "Base de conocimiento propia; mensajes de clientes", "Sí", "No", "Sí", "Clientes PyME y sus contadores",
         "Alto", None, "Líder de atención a clientes", "En producción", "2026-02-10", None, "Aviso de IA al inicio de la conversación"],
        ["IA-03", "Captura automática de CFDI", "Extraer datos de facturas para la contabilidad",
         "Cliente o usuario de IA", "Módulo dentro de otro producto", "Proveedor del software contable", "Visión por computadora",
         "Facturas electrónicas y PDF de clientes", "Sí", "No", "No", "Clientes PyME",
         "Medio", None, "Coordinadora de contabilidad", "En producción", "2026-01-20", None, "Revisión humana por muestreo"],
    ]
    for f, fila in enumerate(ejemplos, 5):
        for c, valor in enumerate(fila, 1):
            ws.cell(row=f, column=c, value=valor)
    ultima = 104
    for f in range(5, ultima + 1):
        ws.cell(row=f, column=14, value=(
            f'=IF(B{f}="","",IF(OR(K{f}="Sí",J{f}="Sí",M{f}="Alto",M{f}="Crítico"),"Sí",'
            f'IF(M{f}="Sin evaluar","Pendiente","Valorar")))'))
        ws.cell(row=f, column=18, value=f'=IF(Q{f}="","",Q{f}+365)')
        ws.cell(row=f, column=18).number_format = "yyyy-mm-dd"
        ws.cell(row=f, column=17).number_format = "yyyy-mm-dd"
        for col in (14, 18):
            ws.cell(row=f, column=col).fill = PatternFill("solid", fgColor=GRIS)
    # Fechas de ejemplo como fechas reales
    from datetime import date
    for f, iso in zip(range(5, 8), ["2026-03-15", "2026-02-10", "2026-01-20"]):
        ws.cell(row=f, column=17, value=date.fromisoformat(iso))
    bordes(ws, f"A5:S{ultima}")
    validacion_lista(ws, f"D5:D{ultima}", listas["Rol"], "Rol según ISO/IEC 22989 (cláusula 4.1).")
    validacion_lista(ws, f"E5:E{ultima}", listas["Origen"])
    validacion_lista(ws, f"G5:G{ultima}", listas["Tipo"])
    for col in "IJK":
        validacion_lista(ws, f"{col}5:{col}{ultima}", listas["SiNo"])
    validacion_lista(ws, f"M5:M{ultima}", listas["Nivel"], "Resultado de la última evaluación de riesgos (6.1.2).")
    validacion_lista(ws, f"P5:P{ultima}", listas["Estado"])
    for nivel, color in NIVEL_COLOR.items():
        ws.conditional_formatting.add(f"M5:M{ultima}", CellIsRule(operator="equal", formula=[f'"{nivel}"'], fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(f"N5:N{ultima}", CellIsRule(operator="equal", formula=['"Sí"'], fill=PatternFill("solid", fgColor="FBCFE8"), font=Font(bold=True)))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:S{ultima}"
    ruta = SALIDA / "inventario-sistemas-ia.xlsx"
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------------- matriz de riesgos

FUENTES_RIESGO = [
    "Complejidad del entorno", "Falta de transparencia o explicabilidad", "Nivel de automatización",
    "Aprendizaje automático (datos y entrenamiento)", "Hardware e infraestructura", "Ciclo de vida del sistema",
    "Madurez tecnológica", "Proveedor o tercero", "Uso indebido previsible", "Seguridad propia de la IA",
    "Privacidad y datos personales", "Otra",
]

ESCALA_CONSECUENCIA = [
    (1, "Insignificante", "Sin pérdida relevante ni atención externa", "Molestia que se corrige en el momento", "Sin efecto perceptible"),
    (2, "Menor", "Pérdida que el área absorbe; queja aislada", "Error que la persona corrige sola en días, sin costo", "Efecto local y aislado"),
    (3, "Moderada", "Reasignar presupuesto; requerimiento de una autoridad; quejas repetidas",
     "Afectación reversible con esfuerzo a derechos u oportunidades", "Afecta de forma reversible a un segmento identificable"),
    (4, "Mayor", "Sanción; pérdida de un cliente clave o de una alianza; prensa negativa",
     "Daño significativo y difícil de revertir: discriminación sistemática, exposición de datos sensibles",
     "Refuerza desigualdades o daña la confianza en un sector"),
    (5, "Severa", "Amenaza la continuidad del negocio o la licencia para operar",
     "Daño grave o irreversible a la vida, la salud, la libertad o el patrimonio",
     "Efecto sistémico o duradero en servicios esenciales o procesos democráticos"),
]
ESCALA_PROBABILIDAD = [
    (1, "Rara", "No se espera en la vida del sistema", "Menos de una vez en cinco años o de 1 en 1 000 000 de decisiones"),
    (2, "Improbable", "Ha ocurrido en otras organizaciones, no en la nuestra", "Una vez cada dos a cinco años"),
    (3, "Posible", "Antecedentes internos aislados", "Una vez al año"),
    (4, "Probable", "Se observa varias veces al año o en las pruebas", "Trimestral o más de 1 en 10 000 decisiones"),
    (5, "Casi segura", "Ya está ocurriendo de forma recurrente", "Mensual o más de 1 en 1 000 decisiones"),
]
# Matriz asimétrica de la guía (cláusula 6): filas = consecuencia 5 → 1; columnas = probabilidad 1 → 5
MATRIZ = {
    5: ["Alto", "Alto", "Crítico", "Crítico", "Crítico"],
    4: ["Medio", "Alto", "Alto", "Crítico", "Crítico"],
    3: ["Bajo", "Medio", "Alto", "Alto", "Crítico"],
    2: ["Bajo", "Bajo", "Medio", "Medio", "Alto"],
    1: ["Bajo", "Bajo", "Bajo", "Medio", "Medio"],
}
ACEPTACION = [
    ("Bajo", "Aceptable con los controles existentes", "Sin plan", "Dueño del sistema de IA", "Anual"),
    ("Medio", "Aceptable con monitoreo; tratar si el costo es razonable", "6 meses, si se trata", "Responsable del SGIA", "Semestral"),
    ("Alto", "No aceptable sin tratamiento", "Plan en 30 días; ejecución en 90", "Comité de IA o dirección designada, informando a la alta dirección", "Trimestral"),
    ("Crítico", "Inaceptable: no se lanza o se suspende la función", "Inmediato", "Sin aceptación permanente; excepción temporal solo con firma de la dirección general", "Mensual"),
]


def clasificar(fila: int, col_c: str, col_p: str) -> str:
    """Fórmula que busca la clasificación en la matriz de la hoja Criterios (B22:F26)."""
    return (f'=IF(OR({col_c}{fila}="",{col_p}{fila}=""),"",'
            f'INDEX(Criterios!$B$22:$F$26,6-{col_c}{fila},{col_p}{fila}))')


def matriz_riesgos() -> Path:
    wb = libro("Metodología y matriz de riesgos de IA")
    hoja_instrucciones(wb, "Metodología y matriz de evaluación de riesgos de IA", [
        "1. Revisa y ajusta la hoja Criterios: escalas de consecuencia (para la organización, para individuos y para la "
        "sociedad), escala de probabilidad, matriz de clasificación y tabla de aceptación. Apruébalas como criterios de "
        "riesgo de IA (cláusula 6.1.1).",
        "2. En Registro de riesgos, identifica cada riesgo por sistema de IA (o grupo de sistemas), con su fuente de riesgo "
        "(inspirada en el Anexo C) y su descripción en formato causa → evento → consecuencia.",
        "3. Califica la consecuencia en las tres dimensiones. La consecuencia que cuenta es la PEOR de las tres (no el "
        "promedio): un riesgo puede ser leve para la organización y grave para las personas.",
        "4. Califica la probabilidad (cuando aplique). La clasificación se obtiene de la matriz asimétrica de la hoja "
        "Criterios, que pesa más la consecuencia que la probabilidad.",
        "5. Define el tratamiento, los controles del Anexo A y los controles propios, y vuelve a calificar el riesgo residual. "
        "Registra quién aprueba la aceptación del riesgo residual según la tabla de aceptación (6.1.3).",
        "6. Usa los resultados de las evaluaciones de impacto (6.1.4) para calificar la consecuencia para individuos y sociedad.",
        "7. Las hojas Mapa inherente y Mapa residual cuentan los riesgos en cada celda de la matriz 5 × 5.",
        "8. Las filas de ejemplo corresponden a la empresa ficticia Monarca Crédito (caso práctico 2) y coinciden con el "
        "ejemplo resuelto de la cláusula 6 en la guía.",
    ])
    listas = hoja_listas(wb, {
        "Fuente": FUENTES_RIESGO,
        "Tratamiento": ["Mitigar", "Evitar", "Transferir o compartir", "Aceptar", "Aprovechar (oportunidad)"],
        "Estado": ["Identificado", "En tratamiento", "Tratado", "Aceptado", "Cerrado"],
        "SiNo": ["Sí", "No", "Temporal"],
    })

    # Criterios
    wc = wb.create_sheet("Criterios", 1)
    titulo_hoja(wc, "Criterios de riesgo de IA", "Ajusta descriptores, matriz y tabla de aceptación a tu contexto y apruébalos (6.1.1).")
    encabezados(wc, 4, [("Nivel", 8), ("Consecuencia", 16), ("Para la organización", 40), ("Para individuos o grupos", 44), ("Para la sociedad", 40)])
    for f, fila in enumerate(ESCALA_CONSECUENCIA, 5):
        for c, v in enumerate(fila, 1):
            wc.cell(row=f, column=c, value=v)
    bordes(wc, "A5:E9")
    encabezados(wc, 12, [("Nivel", 8), ("Probabilidad", 16), ("Descripción", 40), ("Guía de frecuencia (ejemplo)", 44)])
    for f, fila in enumerate(ESCALA_PROBABILIDAD, 13):
        for c, v in enumerate(fila, 1):
            wc.cell(row=f, column=c, value=v)
    bordes(wc, "A13:D17")
    wc["A19"] = "Matriz de clasificación (filas: consecuencia · columnas: probabilidad). Edita los textos si cambias tu apetito de riesgo."
    wc["A19"].font = Font(bold=True, color=INDIGO)
    wc.cell(row=21, column=1, value="C \\ P").font = Font(bold=True)
    for p in range(1, 6):
        c = wc.cell(row=21, column=1 + p, value=p)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for i, consecuencia in enumerate([5, 4, 3, 2, 1]):
        fila = 22 + i
        wc.cell(row=fila, column=1, value=consecuencia).font = Font(bold=True)
        for p in range(5):
            nivel = MATRIZ[consecuencia][p]
            celda = wc.cell(row=fila, column=2 + p, value=nivel)
            celda.fill = PatternFill("solid", fgColor=NIVEL_COLOR[nivel])
            celda.alignment = Alignment(horizontal="center")
            celda.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
    lista_niveles = DataValidation(type="list", formula1='"Bajo,Medio,Alto,Crítico"', allow_blank=False)
    wc.add_data_validation(lista_niveles)
    lista_niveles.add("B22:F26")
    for nivel, color in NIVEL_COLOR.items():
        wc.conditional_formatting.add("B22:F26", CellIsRule(operator="equal", formula=[f'"{nivel}"'], fill=PatternFill("solid", fgColor=color)))
    encabezados(wc, 29, [("Nivel", 12), ("Decisión", 40), ("Plazo", 26), ("Quién acepta el residual", 44), ("Revisión", 40)])
    for f, fila in enumerate(ACEPTACION, 30):
        for c, v in enumerate(fila, 1):
            wc.cell(row=f, column=c, value=v)
        wc.cell(row=f, column=1).fill = PatternFill("solid", fgColor=NIVEL_COLOR[fila[0]])
    bordes(wc, "A30:E33")
    wc["A35"] = ("Líneas rojas (no pasan por la matriz): incumplir deliberadamente una ley aplicable, desplegar un uso de IA "
                 "prohibido en alguna jurisdicción donde operas o usar datos personales para una finalidad no informada.")
    wc["A35"].font = Font(italic=True, color="991B1B")

    # Registro
    ws = wb.create_sheet("Registro de riesgos", 2)
    titulo_hoja(ws, "Registro de riesgos de IA", "Cláusulas 6.1.2, 6.1.3, 8.2 y 8.3")
    cols = [
        ("ID", 8), ("Sistema de IA", 20), ("Fuente de riesgo", 24), ("Descripción (causa → evento → consecuencia)", 50),
        ("Dueño del riesgo", 18), ("Consecuencia organización", 11), ("Consecuencia individuos", 11), ("Consecuencia sociedad", 11),
        ("Consecuencia (peor)", 11), ("Probabilidad", 11), ("Clasificación inherente", 13),
        ("Tratamiento", 14), ("Controles del Anexo A", 24), ("Controles propios o adicionales", 30),
        ("Consecuencia residual", 11), ("Probabilidad residual", 11), ("Clasificación residual", 13),
        ("¿Residual aceptado?", 11), ("Aprobado por", 18), ("Estado", 13), ("Fecha de evaluación", 13),
    ]
    encabezados(ws, 4, cols)
    from datetime import date
    hoy = date(2026, 11, 12)
    ejemplos = [
        ["R-01", "Score Monarca v3", "Aprendizaje automático (datos y entrenamiento)",
         "Variables sustitutas: el código postal se correlaciona con región e ingreso → el modelo podría rechazar de forma "
         "desproporcionada a solicitantes de ciertas entidades → negación de crédito a grupos enteros",
         "Director de Riesgos", 4, 4, 3, None, 3, None, "Mitigar", "A.7.4, A.7.6, A.6.2.4, A.5.4",
         "C-MOD-01 prueba de variables sustitutas en cada reentrenamiento", 4, 1, None, "Sí", "Comité de Modelos", "En tratamiento", hoy],
        ["R-02", "Score Monarca v3", "Ciclo de vida del sistema",
         "Deriva económica: cambios en inflación y empleo → los solicitantes dejan de parecerse a los de entrenamiento → "
         "morosidad y rechazos injustos",
         "Líder de Ciencia de Datos", 4, 3, 2, None, 4, None, "Mitigar", "A.6.2.6, A.6.2.8, A.6.2.5",
         "C-MON-02 interruptor: si el índice de estabilidad poblacional supera el umbral interno, todo pasa a la banda gris",
         3, 2, None, "Sí", "Comité de Modelos", "En tratamiento", hoy],
        ["R-03", "Score Monarca v3", "Falta de transparencia o explicabilidad",
         "Rechazos sin explicación: el rechazo automático muestra un mensaje genérico → los solicitantes no pueden "
         "entender ni impugnar la decisión → quejas ante la autoridad de protección al usuario financiero",
         "Oficial de Cumplimiento", 3, 4, 2, None, 4, None, "Mitigar", "A.8.2, A.8.3, A.6.2.7",
         "C-EXP-01 motivos en lenguaje claro probados con usuarios", 2, 3, None, "Sí", "Comité de Modelos", "En tratamiento", hoy],
        ["R-04", "Score Monarca v3", "Privacidad y datos personales",
         "Datos fuera de finalidad: variables de uso de la app recolectadas para otra finalidad del aviso de privacidad → "
         "se usan en el modelo sin base adecuada → posible incumplimiento de la ley de datos personales",
         "Oficial de Privacidad", 4, 3, 1, None, 2, None, "Evitar", "A.7.3, A.7.5, A.4.3, A.2.3",
         "Retiro de las dos variables", 4, 1, None, "Sí", "Comité de Modelos", "Tratado", hoy],
        ["R-05", "Score Monarca v3", "Nivel de automatización",
         "Sesgo de automatización: por carga de trabajo, los analistas confirman casi siempre la sugerencia del modelo → "
         "la supervisión humana de la banda gris se vuelve ineficaz",
         "Director de Riesgos", 3, 4, 2, None, 3, None, "Mitigar", "A.9.3, A.4.6, A.6.2.6",
         "C-HUM-01: 5 % de la banda gris se presenta sin el score", 4, 2, None, "Temporal", "Comité de Modelos", "En tratamiento", hoy],
    ]
    for f, fila in enumerate(ejemplos, 5):
        for c, valor in enumerate(fila, 1):
            ws.cell(row=f, column=c, value=valor)
    ultima = 154
    for f in range(5, ultima + 1):
        ws.cell(row=f, column=9, value=f'=IF(COUNT(F{f}:H{f})=0,"",MAX(F{f}:H{f}))')
        ws.cell(row=f, column=11, value=clasificar(f, "I", "J"))
        ws.cell(row=f, column=17, value=clasificar(f, "O", "P"))
        ws.cell(row=f, column=21).number_format = "yyyy-mm-dd"
        for col in (9, 11, 17):
            ws.cell(row=f, column=col).fill = PatternFill("solid", fgColor=GRIS)
    bordes(ws, f"A5:U{ultima}")
    validacion_lista(ws, f"C5:C{ultima}", listas["Fuente"], "Fuentes de riesgo inspiradas en el Anexo C.")
    for col in ("F", "G", "H", "J", "O", "P"):
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
        dv.error = "Usa un número del 1 al 5 (hoja Criterios)."
        ws.add_data_validation(dv)
        dv.add(f"{col}5:{col}{ultima}")
    validacion_lista(ws, f"L5:L{ultima}", listas["Tratamiento"])
    validacion_lista(ws, f"R5:R{ultima}", listas["SiNo"])
    validacion_lista(ws, f"T5:T{ultima}", listas["Estado"])
    for rango in (f"K5:K{ultima}", f"Q5:Q{ultima}"):
        for nivel, color in NIVEL_COLOR.items():
            ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=[f'"{nivel}"'], fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(
        f"R5:R{ultima}",
        FormulaRule(formula=['AND(Q5="Crítico",R5="Sí")'], fill=PatternFill("solid", fgColor="FECACA"), font=Font(bold=True)))
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:U{ultima}"

    # Mapas de calor
    for nombre, col_c, col_p in (("Mapa inherente", "I", "J"), ("Mapa residual", "O", "P")):
        wm = wb.create_sheet(nombre)
        titulo_hoja(wm, f"{nombre} de riesgos de IA", "Cada celda cuenta los riesgos del registro con esa consecuencia y probabilidad.")
        wm["B4"] = "Consecuencia ↓ / Probabilidad →"
        wm["B4"].font = Font(bold=True, color=INDIGO)
        wm.column_dimensions["A"].width = 4
        wm.column_dimensions["B"].width = 20
        for i, (n, etiqueta, *_r) in enumerate(ESCALA_PROBABILIDAD):
            c = wm.cell(row=5, column=3 + i, value=f"{n} · {etiqueta}")
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", wrap_text=True)
            wm.column_dimensions[get_column_letter(3 + i)].width = 16
        for j, consecuencia in enumerate([5, 4, 3, 2, 1]):
            fila = 6 + j
            etiqueta = ESCALA_CONSECUENCIA[consecuencia - 1][1]
            wm.cell(row=fila, column=2, value=f"{consecuencia} · {etiqueta}").font = Font(bold=True)
            wm.row_dimensions[fila].height = 42
            for p in range(1, 6):
                celda = wm.cell(row=fila, column=2 + p, value=(
                    f"=COUNTIFS('Registro de riesgos'!${col_c}$5:${col_c}${ultima},{consecuencia},"
                    f"'Registro de riesgos'!${col_p}$5:${col_p}${ultima},{p})"))
                celda.fill = PatternFill("solid", fgColor=NIVEL_COLOR[MATRIZ[consecuencia][p - 1]])
                celda.alignment = Alignment(horizontal="center", vertical="center")
                celda.font = Font(bold=True, size=14)
                celda.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)
        wm["B12"] = "Los colores siguen la matriz por defecto. Si cambias la matriz en Criterios, ajusta también estos colores."
        wm["B12"].font = Font(italic=True, size=9, color="555555")
    ruta = SALIDA / "matriz-riesgos-ia.xlsx"
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------------- SoA

def soa(datos: dict) -> tuple[Path, Path]:
    objetivos = {o["id"]: o for o in datos["objetivos"]}
    wb = libro("Declaración de Aplicabilidad ISO/IEC 42001")
    hoja_instrucciones(wb, "Declaración de Aplicabilidad (SoA) · 38 controles del Anexo A", [
        "1. Para cada control decide si aplica (Sí/No) a partir de tu evaluación de riesgos, tus evaluaciones de impacto y "
        "tus requisitos externos (6.1.3).",
        "2. Justifica TODAS las inclusiones y exclusiones. Una exclusión puede apoyarse en que la evaluación de riesgos no "
        "lo requiere y ningún requisito externo lo exige; escribe el porqué concreto, no frases genéricas.",
        "3. Indica el estado de implementación, los riesgos que trata (IDs del registro) y la evidencia real (ruta o enlace).",
        "4. La columna Validación revisa la consistencia: avisa si falta decisión, justificación, estado o evidencia.",
        "5. Si agregas controles propios fuera del Anexo A, añádelos al final con el prefijo 'P-'.",
        "6. La columna 'Evidencia típica' es orientativa (interpretación del autor). La guía de cada control está en el sitio.",
        "7. Los nombres de los controles son traducciones libres de referencia del autor.",
    ])
    listas = hoja_listas(wb, {
        "Aplica": ["Sí", "No"],
        "Estado": ["No iniciado", "En curso", "Implementado", "No aplica"],
    })
    ws = wb.create_sheet("SoA", 1)
    titulo_hoja(ws, "Declaración de Aplicabilidad · ISO/IEC 42001 Anexo A", "Organización: ____________________   Versión: ____   Aprobada por: ____________________   Fecha: ________")
    cols = [
        ("Control", 10), ("Nombre (traducción libre)", 34), ("Objetivo", 22), ("Aplica a (orientativo)", 20),
        ("¿Aplica?", 9), ("Justificación de inclusión o exclusión", 44), ("Estado de implementación", 15),
        ("Riesgos que trata (IDs)", 16), ("Evidencia típica (orientativa)", 40), ("Evidencia real (ruta o enlace)", 30),
        ("Responsable", 18), ("ISO 27001:2022 relacionado", 14), ("Validación", 26),
    ]
    encabezados(ws, 4, cols)
    filas_csv = []
    for f, c in enumerate(datos["controles"], 5):
        o = objetivos[c["objetivo"]]
        roles = " / ".join(ROLES[r] for r in c["roles"])
        mapeo = ", ".join(c["iso27001"]) if c["iso27001"] else "—"
        valores = [c["id"], c["nombre"], f'{o["id"]} {o["nombre"]}', roles, None, None, None, None, c["evidencia"], None, None, mapeo]
        for col, v in enumerate(valores, 1):
            ws.cell(row=f, column=col, value=v)
        ws.cell(row=f, column=13, value=(
            f'=IF(E{f}="","Pendiente: decide si aplica",'
            f'IF(F{f}="","Falta justificación",'
            f'IF(AND(E{f}="Sí",G{f}=""),"Falta estado",'
            f'IF(AND(E{f}="Sí",G{f}="No aplica"),"Inconsistente: aplica pero estado No aplica",'
            f'IF(AND(E{f}="No",G{f}<>"",G{f}<>"No aplica"),"Inconsistente: excluido con estado",'
            f'IF(AND(E{f}="Sí",G{f}="Implementado",J{f}=""),"Falta evidencia","OK"))))))'))
        ws.cell(row=f, column=13).fill = PatternFill("solid", fgColor=GRIS)
        filas_csv.append([c["id"], c["nombre"], o["id"], o["nombre"], roles, c["esfuerzo"], NOVEDAD[c["novedad"]], mapeo,
                          c["resumen"], c["evidencia"], "", "", "", "", ""])
    ultima = 4 + len(datos["controles"])
    bordes(ws, f"A5:M{ultima}")
    validacion_lista(ws, f"E5:E{ultima}", listas["Aplica"])
    validacion_lista(ws, f"G5:G{ultima}", listas["Estado"])
    ws.conditional_formatting.add(f"M5:M{ultima}", CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="BBF7D0")))
    ws.conditional_formatting.add(f"M5:M{ultima}", CellIsRule(operator="notEqual", formula=['"OK"'], fill=PatternFill("solid", fgColor="FEF08A")))
    ws.conditional_formatting.add(f"E5:E{ultima}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="E5E7EB")))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:M{ultima}"

    # Resumen
    wr = wb.create_sheet("Resumen", 2)
    titulo_hoja(wr, "Resumen de la Declaración de Aplicabilidad")
    wr.column_dimensions["A"].width = 46
    wr.column_dimensions["B"].width = 14
    indicadores = [
        ("Controles del Anexo A", f"=COUNTA(SoA!A5:A{ultima})"),
        ("Controles que aplican", f'=COUNTIF(SoA!E5:E{ultima},"Sí")'),
        ("Controles excluidos", f'=COUNTIF(SoA!E5:E{ultima},"No")'),
        ("Pendientes de decisión", f'=COUNTBLANK(SoA!E5:E{ultima})'),
        ("Implementados", f'=COUNTIFS(SoA!E5:E{ultima},"Sí",SoA!G5:G{ultima},"Implementado")'),
        ("En curso", f'=COUNTIFS(SoA!E5:E{ultima},"Sí",SoA!G5:G{ultima},"En curso")'),
        ("No iniciados", f'=COUNTIFS(SoA!E5:E{ultima},"Sí",SoA!G5:G{ultima},"No iniciado")'),
        ("Avance de implementación (implementados / aplican)", '=IF(B5=0,"",B8/B5)'),
        ("Filas con validación distinta de OK", f'=COUNTIF(SoA!M5:M{ultima},"<>OK")'),
    ]
    encabezados(wr, 3, [("Indicador", 46), ("Valor", 14)])
    for f, (nombre, formula) in enumerate(indicadores, 4):
        wr.cell(row=f, column=1, value=nombre)
        wr.cell(row=f, column=2, value=formula)
    wr["B11"].number_format = "0%"
    bordes(wr, "A4:B12")
    encabezados(wr, 15, [("Objetivo", 46), ("Aplican", 14), ("Implementados", 14)])
    wr.column_dimensions["C"].width = 14
    for f, o in enumerate(datos["objetivos"], 16):
        wr.cell(row=f, column=1, value=f'{o["id"]} {o["nombre"]}')
        wr.cell(row=f, column=2, value=f'=COUNTIFS(SoA!C5:C{ultima},"{o["id"]} *",SoA!E5:E{ultima},"Sí")')
        wr.cell(row=f, column=3, value=f'=COUNTIFS(SoA!C5:C{ultima},"{o["id"]} *",SoA!E5:E{ultima},"Sí",SoA!G5:G{ultima},"Implementado")')
    bordes(wr, f"A16:C{15 + len(datos['objetivos'])}")
    ruta = SALIDA / "declaracion-de-aplicabilidad.xlsx"
    wb.save(ruta)

    ruta_csv = SALIDA / "declaracion-de-aplicabilidad.csv"
    with ruta_csv.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["control", "nombre_traduccion_libre", "objetivo", "nombre_objetivo", "aplica_a_orientativo",
                    "esfuerzo_estimado", "frente_a_iso27001", "iso27001_2022_relacionado", "resumen", "evidencia_tipica",
                    "aplica_si_no", "justificacion", "estado_implementacion", "evidencia_real", "responsable"])
        w.writerows(filas_csv)
    return ruta, ruta_csv


# ----------------------------------------------------------------------------- incidentes

def incidentes() -> Path:
    wb = libro("Registro de incidentes de IA")
    hoja_instrucciones(wb, "Registro de incidentes de IA", [
        "1. Registra todo evento en que un sistema de IA haya causado o pudo causar un daño, un incumplimiento o un "
        "resultado inaceptable: errores graves, alucinaciones con impacto, trato injusto, ataques, fugas de datos, uso indebido.",
        "2. Clasifica el tipo y la severidad. Decide si hay que comunicarlo a usuarios (A.8.4) y si hay obligación de "
        "notificar a una autoridad (A.8.5) según tu contexto legal.",
        "3. Si el incidente revela una falla del SGIA, abre una no conformidad y una acción correctiva (10.2) y anota su ID.",
        "4. Las columnas en gris se calculan solas (días abiertos y alerta de seguimiento).",
        "5. Revisa la hoja Resumen en la revisión por la dirección (9.3).",
        "6. Las filas de ejemplo corresponden a la empresa ficticia Conversa Labs (caso práctico 3).",
    ])
    listas = hoja_listas(wb, {
        "Tipo": ["Desempeño o error del modelo", "Alucinación con impacto", "Sesgo o trato injusto", "Seguridad propia de la IA",
                 "Privacidad o datos personales", "Uso indebido o fuera del uso previsto", "Falla del proveedor", "Otro"],
        "Severidad": ["Baja", "Media", "Alta", "Crítica"],
        "SiNo": ["Sí", "No", "Por determinar"],
        "Estado": ["Abierto", "En análisis", "Contenido", "Cerrado"],
        "Origen": ["Monitoreo automático", "Reporte de usuario", "Reporte de cliente", "Reporte interno (A.3.3)",
                   "Proveedor", "Auditoría", "Otro"],
    })
    ws = wb.create_sheet("Registro", 1)
    titulo_hoja(ws, "Registro de incidentes de IA", "Controles relacionados: A.6.2.6, A.6.2.8, A.8.3, A.8.4, A.8.5, A.10.3 · Cláusula 10.2")
    cols = [
        ("ID", 9), ("Fecha de detección", 13), ("Sistema de IA", 22), ("Origen del reporte", 18), ("Descripción", 44),
        ("Tipo", 22), ("Severidad", 11), ("Personas afectadas (aprox.)", 12), ("¿Afecta a personas externas?", 12),
        ("¿Comunicar a usuarios? (A.8.4)", 13), ("¿Notificar a autoridad?", 13), ("Contención aplicada", 32),
        ("Causa raíz", 32), ("ID de no conformidad / acción correctiva", 16), ("Responsable", 18), ("Estado", 12),
        ("Fecha de cierre", 13), ("Días abiertos", 10), ("Alerta", 22), ("Lecciones aprendidas", 32),
    ]
    encabezados(ws, 4, cols)
    from datetime import date
    ejemplos = [
        ["INC-2026-014", date(2026, 6, 3), "Conversa · asistente de aseguradora", "Reporte de cliente",
         "El asistente afirmó que una póliza cubría un procedimiento que no cubre", "Alucinación con impacto", "Alta", 37, "Sí",
         "Sí", "Por determinar", "Se desactivó la intención afectada y se activó respuesta con traspaso a humano",
         "Documento desactualizado en la base de conocimiento y verificación insuficiente de fuentes", "AC-2026-009",
         "Responsable de Confianza y Seguridad", "Cerrado", date(2026, 6, 24), None, None,
         "Agregar prueba de regresión con preguntas de coberturas antes de cada actualización de la base"],
        ["INC-2026-015", date(2026, 7, 11), "Conversa · asistente universitario", "Monitoreo automático",
         "Intento de inyección de instrucciones para obtener el texto del sistema", "Seguridad propia de la IA", "Media", 0, "No",
         "No", "No", "Bloqueo por filtro y revisión de registros", None, None, "CTO", "En análisis", None, None, None, None],
    ]
    for f, fila in enumerate(ejemplos, 5):
        for c, valor in enumerate(fila, 1):
            ws.cell(row=f, column=c, value=valor)
    ultima = 204
    for f in range(5, ultima + 1):
        ws.cell(row=f, column=18, value=f'=IF(OR(B{f}="",Q{f}=""),"",Q{f}-B{f})')
        ws.cell(row=f, column=19, value=(
            f'=IF(A{f}="","",IF(AND(P{f}<>"Cerrado",OR(G{f}="Alta",G{f}="Crítica"),J{f}=""),'
            f'"Decidir comunicación a usuarios",IF(AND(P{f}="Cerrado",M{f}=""),"Cerrado sin causa raíz","")))'))
        for col in (2, 17):
            ws.cell(row=f, column=col).number_format = "yyyy-mm-dd"
        for col in (18, 19):
            ws.cell(row=f, column=col).fill = PatternFill("solid", fgColor=GRIS)
    bordes(ws, f"A5:T{ultima}")
    validacion_lista(ws, f"D5:D{ultima}", listas["Origen"])
    validacion_lista(ws, f"F5:F{ultima}", listas["Tipo"])
    validacion_lista(ws, f"G5:G{ultima}", listas["Severidad"])
    for col in "IJK":
        validacion_lista(ws, f"{col}5:{col}{ultima}", listas["SiNo"])
    validacion_lista(ws, f"P5:P{ultima}", listas["Estado"])
    sev = {"Baja": "BBF7D0", "Media": "FEF08A", "Alta": "FED7AA", "Crítica": "FECACA"}
    for nivel, color in sev.items():
        ws.conditional_formatting.add(f"G5:G{ultima}", CellIsRule(operator="equal", formula=[f'"{nivel}"'], fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add(f"S5:S{ultima}", CellIsRule(operator="notEqual", formula=['""'], fill=PatternFill("solid", fgColor="FECACA")))
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:T{ultima}"

    wr = wb.create_sheet("Resumen", 2)
    titulo_hoja(wr, "Resumen de incidentes de IA")
    encabezados(wr, 3, [("Tipo de incidente", 36), ("Total", 10), ("Abiertos", 10), ("Alta o crítica", 14)])
    tipos = ["Desempeño o error del modelo", "Alucinación con impacto", "Sesgo o trato injusto", "Seguridad propia de la IA",
             "Privacidad o datos personales", "Uso indebido o fuera del uso previsto", "Falla del proveedor", "Otro"]
    for f, t in enumerate(tipos, 4):
        wr.cell(row=f, column=1, value=t)
        wr.cell(row=f, column=2, value=f'=COUNTIF(Registro!F5:F{ultima},A{f})')
        wr.cell(row=f, column=3, value=f'=COUNTIFS(Registro!F5:F{ultima},A{f},Registro!P5:P{ultima},"<>Cerrado")')
        wr.cell(row=f, column=4, value=f'=COUNTIFS(Registro!F5:F{ultima},A{f},Registro!G5:G{ultima},"Alta")+COUNTIFS(Registro!F5:F{ultima},A{f},Registro!G5:G{ultima},"Crítica")')
    total = 4 + len(tipos)
    wr.cell(row=total, column=1, value="Total").font = Font(bold=True)
    for col in "BCD":
        wr[f"{col}{total}"] = f"=SUM({col}4:{col}{total - 1})"
        wr[f"{col}{total}"].font = Font(bold=True)
    wr.cell(row=total + 2, column=1, value="Promedio de días para cerrar")
    wr.cell(row=total + 2, column=2, value=f'=IFERROR(AVERAGE(Registro!R5:R{ultima}),"")')
    wr.cell(row=total + 2, column=2).number_format = "0.0"
    bordes(wr, f"A4:D{total}")
    ruta = SALIDA / "registro-incidentes-ia.xlsx"
    wb.save(ruta)
    return ruta


def main() -> None:
    SALIDA.mkdir(exist_ok=True)
    datos = yaml.safe_load(DATOS.read_text(encoding="utf-8"))
    rutas = [inventario(), matriz_riesgos(), *soa(datos), incidentes()]
    for r in rutas:
        print(f"OK -> {r.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
